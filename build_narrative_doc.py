from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / '信息学网页开场叙事方案.md'
TARGET = ROOT / '信息学网页开场叙事方案.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.75), Inches(.75)
sec.left_margin, sec.right_margin = Inches(.8), Inches(.8)

for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.25
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].paragraph_format.space_after = Pt(16)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(6)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(11)
doc.core_properties.title = '信息学网页开场叙事方案'
doc.core_properties.subject = '学生活动网页内容与讲解衔接'
doc.core_properties.author = ''

def add_text(p, text):
    pos = 0
    for m in re.finditer(r'\[([^\]]+)\]\((https?://[^)]+)\)', text):
        p.add_run(text[pos:m.start()])
        rel_id = p.part.relate_to(m.group(2), 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        h = OxmlElement('w:hyperlink')
        h.set(qn('r:id'), rel_id)
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '245A81'); props.append(color)
        run.append(props)
        t = OxmlElement('w:t'); t.text = m.group(1); run.append(t)
        h.append(run); p._p.append(h)
        pos = m.end()
    p.add_run(text[pos:])

def add_table(lines):
    rows = [[c.strip() for c in line.strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.match(r'^:?-+:?$', c) for c in r)]
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.15, 2.2, 3.55] if rows[0][0] == '学科' else [1.15, 2.6, 3.15]
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for idx, row in enumerate(rows):
        cells = table.rows[0].cells if idx == 0 else table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text); r.font.size = Pt(10)
            r.font.bold = idx == 0
            if idx == 0: r.font.color.rgb = RGBColor(255, 255, 255)
            tcpr = cells[i]._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), '244A61' if idx == 0 else ('EFF4F7' if idx % 2 == 0 else 'FFFFFF')); tcpr.append(shd)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                b = OxmlElement('w:' + edge); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9'); borders.append(b)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                mar = OxmlElement('w:' + edge); mar.set(qn('w:w'), '90'); mar.set(qn('w:type'), 'dxa'); margins.append(mar)
            tcpr.append(margins)
        if idx == 0:
            trpr = table.rows[0]._tr.get_or_add_trPr()
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

lines = SOURCE.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line: i += 1; continue
    if line == '<!-- pagebreak -->':
        doc.add_page_break()
    elif line.startswith('|'):
        table_lines = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            table_lines.append(lines[i].strip()); i += 1
        add_table(table_lines)
        continue
    elif line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_heading(line[3:], 1)
    elif line.startswith('### '):
        doc.add_heading(line[4:], 2)
    else:
        p = doc.add_paragraph()
        add_text(p, line)
    i += 1
doc.save(TARGET)
print(TARGET)
